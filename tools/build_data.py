# -*- coding: utf-8 -*-
"""Genera data/alimentos.js para NutricionAPP a partir de la base de datos.

Fuentes posibles (en la carpeta padre de la app):
  alimentos_nutrientes_mexico.csv   -> filas (118 x 75)
  alimentos_nutrientes_mexico.xlsx  -> hoja "Base de datos"
  alimentos_nutrientes_mexico.db    -> tabla "alimentos"

El diccionario (etiquetas, unidades y descripciones) se toma de la tabla
"diccionario" del .db; si no existe, se usan las cabeceras del archivo.

El resultado se carga con <script src="data/alimentos.js">, sin fetch, para que
la app funcione abriendo index.html directamente (file://).

Uso:
    py tools/build_data.py                 # detecta la mejor fuente disponible
    py tools/build_data.py --source csv    # fuerza CSV
    py tools/build_data.py --source xlsx   # fuerza Excel
    py tools/build_data.py --source db     # fuerza SQLite
"""
import csv
import json
import os
import re
import shutil
import sqlite3
import sys
from datetime import date

AQUI = os.path.dirname(os.path.abspath(__file__))
APP = os.path.dirname(AQUI)
RAIZ = os.path.dirname(APP)

CSV = os.path.join(RAIZ, "alimentos_nutrientes_mexico.csv")
XLSX = os.path.join(RAIZ, "alimentos_nutrientes_mexico.xlsx")
DB = os.path.join(RAIZ, "alimentos_nutrientes_mexico.db")
ICONOS_ORIGEN = os.path.join(RAIZ, "iconos")
SALIDA = os.path.join(APP, "data", "alimentos.js")
DESTINO_ICONOS = os.path.join(APP, "iconos")

# Orden exacto de las 75 columnas del CSV/XLSX.
KEYS = [
    "id", "nombre", "grupo", "categoria", "porcion", "unidad", "equiv",
    "fuente", "notas",
    "water", "kcal", "protein", "fat", "sfa", "mufa", "pufa", "insat",
    "trans", "cholesterol", "carb", "sugars", "fiber", "starch", "ash",
    "la", "ala", "epa", "dha", "dpa", "ara", "n3", "n6", "ratio",
    "va", "retinol", "beta_car", "alpha_car", "beta_cryp", "lycopene",
    "lutein", "vc", "vd", "ve", "vk", "b1", "b2", "b3", "b5", "b6",
    "biotin", "folate", "b12", "choline", "betaine",
    "ca", "fe", "mg", "p", "k", "na", "zn", "cu", "mn", "se", "iodine",
    "chromium", "molybdenum", "fluoride",
    "phytosterols", "beta_sito", "caffeine", "theobromine",
    "antiox", "otros", "icono",
]

SECCION = {}
for k in ("id", "nombre", "grupo", "categoria", "porcion", "unidad",
          "equiv", "fuente", "notas", "icono"):
    SECCION[k] = "ident"
for k in ("water", "kcal", "protein", "fat", "sfa", "mufa", "pufa", "insat",
          "trans", "cholesterol", "carb", "sugars", "fiber", "starch", "ash"):
    SECCION[k] = "macro"
for k in ("la", "ala", "epa", "dha", "dpa", "ara", "n3", "n6", "ratio"):
    SECCION[k] = "grasas"
for k in ("va", "retinol", "beta_car", "alpha_car", "beta_cryp", "lycopene",
          "lutein", "vc", "vd", "ve", "vk", "b1", "b2", "b3", "b5", "b6",
          "biotin", "folate", "b12", "choline", "betaine"):
    SECCION[k] = "vit"
for k in ("ca", "fe", "mg", "p", "k", "na", "zn", "cu", "mn", "se", "iodine",
          "chromium", "molybdenum", "fluoride"):
    SECCION[k] = "min"
for k in ("phytosterols", "beta_sito", "caffeine", "theobromine"):
    SECCION[k] = "fito"
for k in ("antiox", "otros"):
    SECCION[k] = "texto"

TEXTO = {"antiox", "otros", "id", "nombre", "grupo", "categoria", "porcion",
         "unidad", "equiv", "fuente", "notas", "icono"}

# Metadatos que se guardan como campos propios, no dentro de "nutrientes".
META_FIJOS = ("id", "nombre", "grupo", "categoria", "porcion", "unidad",
              "equiv", "fuente", "notas", "icono")


def num(valor):
    if valor is None:
        return None
    valor = str(valor).strip()
    if valor == "":
        return None
    try:
        f = float(valor)
    except ValueError:
        return None
    return round(f, 4)


def txt(valor):
    if valor is None:
        return ""
    if isinstance(valor, float) and valor.is_integer():
        valor = int(valor)
    return str(valor).strip()


# --------------------------------------------------------------------- fuentes

def leer_csv():
    with open(CSV, encoding="utf-8-sig", newline="") as f:
        lector = csv.reader(f)
        cabeceras = next(lector)
        filas = [dict(zip(cabeceras, vals)) for vals in lector if any(vals)]
    return cabeceras, filas, "alimentos_nutrientes_mexico.csv"


def leer_xlsx():
    from openpyxl import load_workbook
    wb = load_workbook(XLSX, data_only=True, read_only=True)
    if "Base de datos" not in wb.sheetnames:
        wb.close()
        raise SystemExit('El Excel no tiene la hoja "Base de datos".')
    ws = wb["Base de datos"]
    it = ws.iter_rows(values_only=True)
    cabeceras = [txt(c) for c in next(it)]
    filas = []
    for vals in it:
        if not vals or not any(txt(v) for v in vals):
            continue
        filas.append({cabeceras[i]: vals[i] for i in range(len(cabeceras))})
    wb.close()
    return cabeceras, filas, "alimentos_nutrientes_mexico.xlsx (hoja Base de datos)"


def leer_db():
    con = sqlite3.connect(DB)
    cur = con.execute("SELECT * FROM alimentos")
    cabeceras = [d[0] for d in cur.description]
    filas = [dict(zip(cabeceras, r)) for r in cur.fetchall()]
    con.close()
    return cabeceras, filas, "alimentos_nutrientes_mexico.db (tabla alimentos)"


def elegir_fuente(preferida):
    opciones = [
        ("csv", CSV, leer_csv),
        ("xlsx", XLSX, leer_xlsx),
        ("db", DB, leer_db),
    ]
    if preferida and preferida != "auto":
        for nombre, ruta, lector in opciones:
            if nombre == preferida:
                if not os.path.exists(ruta):
                    raise SystemExit(f"No existe la fuente {nombre}: {ruta}")
                return lector()
        raise SystemExit(f"Fuente desconocida: {preferida}")

    for nombre, ruta, lector in opciones:
        if os.path.exists(ruta):
            return lector()
    raise SystemExit("No se encontró ninguna base: CSV, XLSX ni DB.")


def leer_diccionario():
    """key -> (etiqueta, unidad, descripcion) desde la tabla diccionario."""
    dic = {}
    if not os.path.exists(DB):
        return dic
    con = sqlite3.connect(DB)
    try:
        for k, et, un, de in con.execute(
                "SELECT columna, etiqueta, unidad, descripcion FROM diccionario"):
            dic[k] = (txt(et) or k, txt(un), txt(de))
    except sqlite3.Error:
        pass
    con.close()
    return dic


def main():
    preferida = "auto"
    if "--source" in sys.argv:
        i = sys.argv.index("--source")
        if i + 1 < len(sys.argv):
            preferida = sys.argv[i + 1].lower()

    cabeceras, filas, fuente = elegir_fuente(preferida)

    if len(cabeceras) != len(KEYS):
        raise SystemExit(f"Se esperaban {len(KEYS)} columnas y hay {len(cabeceras)}")
    if not filas:
        raise SystemExit("La fuente no tiene filas.")

    dic = leer_diccionario()

    cols = []
    for i, cab in enumerate(cabeceras):
        k = KEYS[i]
        et, un, de = dic.get(k, ("", "", ""))
        if not et:
            et = cab
        # si el diccionario no trae unidad, se intenta leer del encabezado
        if not un:
            m = re.search(r"\(([^()]*)\)\s*$", cab)
            if m:
                un = m.group(1)
        # el diccionario guarda la etiqueta con la unidad entre paréntesis;
        # la unidad ya viaja aparte, así que se quita del rótulo.
        et = re.sub(r"\s*\([^()]*\)\s*$", "", et).strip() or k
        # decimales con los que la tabla original escribe la columna
        dec = 0
        if k not in TEXTO:
            for fila in filas:
                celda = txt(fila.get(cab))
                if "." in celda:
                    dec = max(dec, len(celda.split(".")[1]))
        cols.append({
            "k": k,
            "label": et,
            "unidad": un,
            "desc": de,
            "sec": SECCION[k],
            "tipo": "texto" if k in TEXTO else "num",
            "dec": min(dec, 4),
            "columna": cab,
        })

    alimentos = []
    for fila in filas:
        reg = {
            "id": txt(fila.get(cabeceras[0])),
            "nombre": txt(fila.get(cabeceras[1])),
            "grupo": txt(fila.get(cabeceras[2])),
            "categoria": txt(fila.get(cabeceras[3])),
            "porcion": num(fila.get(cabeceras[4])) or 100,
            "unidad": txt(fila.get(cabeceras[5])),
            "equiv": txt(fila.get(cabeceras[6])),
            "fuente": txt(fila.get(cabeceras[7])),
            "notas": txt(fila.get(cabeceras[8])),
            "icono": txt(fila.get(cabeceras[74])),
            "nutrientes": {},
            "texto": {},
        }
        for i, cab in enumerate(cabeceras):
            k = KEYS[i]
            if k in META_FIJOS:
                continue
            if k in ("antiox", "otros"):
                reg["texto"][k] = txt(fila.get(cab))
            else:
                reg["nutrientes"][k] = num(fila.get(cab))
        alimentos.append(reg)

    categorias = sorted({a["categoria"] for a in alimentos})
    grupos = sorted({a["grupo"] for a in alimentos})

    doc = {
        "meta": {
            "app": "NutricionAPP",
            "version": "1.0",
            "generado": date.today().isoformat(),
            "fuente": fuente,
            "filas": len(alimentos),
            "columnas": len(cabeceras),
        },
        "categorias": categorias,
        "grupos": grupos,
        "cols": cols,
        "alimentos": alimentos,
    }

    os.makedirs(os.path.dirname(SALIDA), exist_ok=True)
    with open(SALIDA, "w", encoding="utf-8") as f:
        f.write("/* Generado por tools/build_data.py. No editar a mano. */\n")
        f.write("window.DB = ")
        json.dump(doc, f, ensure_ascii=False, separators=(",", ":"))
        f.write(";\n")

    # copiar iconos para que la app sea autocontenida
    if os.path.isdir(ICONOS_ORIGEN):
        os.makedirs(DESTINO_ICONOS, exist_ok=True)
        copiados = 0
        for nombre in os.listdir(ICONOS_ORIGEN):
            if nombre.endswith(".svg"):
                shutil.copy2(os.path.join(ICONOS_ORIGEN, nombre),
                             os.path.join(DESTINO_ICONOS, nombre))
                copiados += 1
        print("Iconos copiados:", copiados)

    kb = os.path.getsize(SALIDA) / 1024
    print(f"Fuente: {fuente}")
    print(f"data/alimentos.js: {doc['meta']['filas']} alimentos, "
          f"{doc['meta']['columnas']} columnas, {len(categorias)} categorías, "
          f"{kb:.1f} KB")


if __name__ == "__main__":
    main()
